"""Conversões com planilhas do Excel (openpyxl)."""

import csv
import json
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .base import Conversion, Option, unique
from .data import DELIMITER, read_csv

ALL_SHEETS = Option("all_sheets", "Todas as abas (um arquivo por aba)", False, bool)


def _write_rows(rows: list[list], dest: Path) -> None:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(dest)


# Até 10 dígitos: números maiores costumam ser identificadores (CPF, telefone)
NUMBER = re.compile(r"-?(?:0|[1-9]\d{0,9})(?:\.\d+)?")


def _typed(value: str):
    """Número vira número no Excel; o resto (ex.: "01234", CPFs) continua texto."""
    if NUMBER.fullmatch(value):
        return float(value) if "." in value else int(value)
    return value


def csv_to_xlsx(src: Path, out_dir: Path) -> list[Path]:
    dest = unique(out_dir / f"{src.stem}.xlsx")
    header, rows = read_csv(src)
    _write_rows([header, *([_typed(v) for v in row] for row in rows)], dest)
    return [dest]


def json_to_xlsx(src: Path, out_dir: Path) -> list[Path]:
    dest = unique(out_dir / f"{src.stem}.xlsx")
    data = json.loads(src.read_text(encoding="utf-8-sig"))
    if isinstance(data, dict):
        data = [data]
    if not isinstance(data, list) or not all(isinstance(r, dict) for r in data):
        raise ValueError("o JSON precisa ser um objeto ou uma lista de objetos")
    fields = list(dict.fromkeys(k for row in data for k in row))
    rows = [[_cell(row.get(k)) for k in fields] for row in data]
    _write_rows([fields, *rows], dest)
    return [dest]


def _cell(value):
    return json.dumps(value, ensure_ascii=False) if isinstance(value, (dict, list)) else value


def _sheets(src: Path, all_sheets: bool):
    wb = load_workbook(src, read_only=True, data_only=True)
    sheets = wb.worksheets if all_sheets else [wb.active]
    for ws in sheets:
        suffix = f"_{ws.title}" if all_sheets else ""
        yield suffix, [["" if v is None else v for v in row] for row in ws.iter_rows(values_only=True)]
    wb.close()


def xlsx_to_csv(src: Path, out_dir: Path, delimiter: str = ",", all_sheets: bool = False) -> list[Path]:
    written = []
    for suffix, rows in _sheets(src, all_sheets):
        dest = unique(out_dir / f"{src.stem}{suffix}.csv")
        with dest.open("w", encoding="utf-8-sig", newline="") as f:
            csv.writer(f, delimiter=delimiter or ",").writerows(rows)
        written.append(dest)
    return written


def xlsx_to_json(src: Path, out_dir: Path, all_sheets: bool = False) -> list[Path]:
    written = []
    for suffix, rows in _sheets(src, all_sheets):
        dest = unique(out_dir / f"{src.stem}{suffix}.json")
        header, body = (rows[0], rows[1:]) if rows else ([], [])
        records = [dict(zip(map(str, header), row)) for row in body]
        dest.write_text(json.dumps(records, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        written.append(dest)
    return written


CONVERSIONS = [
    Conversion("csv", "xlsx", csv_to_xlsx),
    Conversion("json", "xlsx", json_to_xlsx),
    Conversion("xlsx", "csv", xlsx_to_csv, [DELIMITER, ALL_SHEETS]),
    Conversion("xlsx", "json", xlsx_to_json, [ALL_SHEETS]),
]
